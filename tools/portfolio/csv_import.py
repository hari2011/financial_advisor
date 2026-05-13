"""Smart CSV / Excel importer with broker auto-detection.

Detects the broker format from column headers and maps rows to Holding / MFHolding.
Supports: Zerodha Console, Groww, Angel One, ICICI Direct, 5Paisa, generic.
"""
from __future__ import annotations

import csv
import io
import logging
import re
from pathlib import Path

from tools.portfolio.models import Holding, MFHolding

logger = logging.getLogger("financegpt.portfolio.csv_import")


# ── Column signature → broker mapping ──
_BROKER_SIGNATURES: dict[str, dict] = {
    "zerodha": {
        "detect": ["symbol", "isin", "quantity available", "average price"],
        "symbol": "symbol",
        "isin": "isin",
        "name": "symbol",
        "quantity": "quantity available",
        "avg_price": "average price",
        "current_price": "previous closing price",
        "exchange": "exchange",
    },
    "zerodha_console": {
        "detect": ["instrument", "qty.", "avg. cost"],
        "symbol": "instrument",
        "isin": "isin",
        "name": "instrument",
        "quantity": "qty.",
        "avg_price": "avg. cost",
        "current_price": "ltp",
        "exchange": "exchange",
    },
    "groww": {
        "detect": ["symbol", "company name", "quantity", "avg. cost"],
        "symbol": "symbol",
        "isin": "isin",
        "name": "company name",
        "quantity": "quantity",
        "avg_price": "avg. cost",
        "current_price": "current price",
    },
    "angelone": {
        "detect": ["symbol", "company", "qty", "avg price"],
        "symbol": "symbol",
        "isin": "isin",
        "name": "company",
        "quantity": "qty",
        "avg_price": "avg price",
        "current_price": "ltp",
    },
    "icici_direct": {
        "detect": ["stock code", "company name", "quantity", "average price"],
        "symbol": "stock code",
        "isin": "isin code",
        "name": "company name",
        "quantity": "quantity",
        "avg_price": "average price",
        "current_price": "current market price",
    },
    "fivepaisa": {
        "detect": ["scrip name", "qty", "avg rate"],
        "symbol": "scrip name",
        "isin": "isin",
        "name": "scrip name",
        "quantity": "qty",
        "avg_price": "avg rate",
        "current_price": "current price",
    },
    "hdfc_sec": {
        "detect": ["script name", "holding qty", "avg. buy price"],
        "symbol": "script name",
        "isin": "isin",
        "name": "script name",
        "quantity": "holding qty",
        "avg_price": "avg. buy price",
        "current_price": "current price",
    },
    "motilal": {
        "detect": ["scrip name", "isin", "qty", "avg cost"],
        "symbol": "scrip name",
        "isin": "isin",
        "name": "scrip name",
        "quantity": "qty",
        "avg_price": "avg cost",
        "current_price": "ltp",
    },
}

# Mutual fund CSV signatures
_MF_SIGNATURES: dict[str, dict] = {
    "zerodha_mf": {
        "detect": ["symbol", "isin", "quantity available", "average price", "instrument type"],
        "scheme_name": "symbol",
        "isin": "isin",
        "folio": "",
        "units": "quantity available",
        "nav": "previous closing price",
        "invested": "",
        "current": "",
        "avg_nav": "average price",
        "instrument_type": "instrument type",
    },
    "groww_mf": {
        "detect": ["scheme name", "folio", "units", "nav"],
        "scheme_name": "scheme name",
        "folio": "folio",
        "units": "units",
        "nav": "nav",
        "invested": "invested value",
        "current": "current value",
    },
    "generic_mf": {
        "detect": ["scheme", "units"],
        "scheme_name": "scheme",
        "folio": "folio",
        "units": "units",
        "nav": "nav",
        "invested": "invested",
        "current": "current value",
    },
}


def _clean_num(val: str) -> float:
    """Parse a number string, removing commas and currency symbols."""
    if not val:
        return 0.0
    cleaned = re.sub(r"[₹$,\s]", "", val.strip())
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def _normalise_headers(row: list[str]) -> list[str]:
    """Lowercase and strip headers."""
    return [h.strip().lower() for h in row]


def _detect_broker(headers: list[str]) -> tuple[str, dict]:
    """Auto-detect broker from column headers. Returns (broker_id, mapping)."""
    normalised = set(headers)
    best_match = ("generic", {})
    best_score = 0

    for broker_id, config in _BROKER_SIGNATURES.items():
        detect_cols = [c.lower() for c in config["detect"]]
        score = sum(1 for c in detect_cols if c in normalised)
        if score > best_score:
            best_score = score
            best_match = (broker_id, config)

    if best_score == 0:
        # Fallback: generic detection
        return _detect_generic(headers)

    return best_match


def _detect_generic(headers: list[str]) -> tuple[str, dict]:
    """Build a mapping for an unknown CSV format by guessing column roles."""
    mapping: dict = {"detect": []}
    header_set = set(headers)

    # Symbol
    for candidate in ["symbol", "ticker", "scrip", "stock", "instrument", "code"]:
        for h in headers:
            if candidate in h:
                mapping["symbol"] = h
                break
        if "symbol" in mapping:
            break

    # Name
    for candidate in ["name", "company", "scrip name", "stock name", "description"]:
        for h in headers:
            if candidate in h:
                mapping["name"] = h
                break
        if "name" in mapping:
            break

    # Quantity
    for candidate in ["qty", "quantity", "shares", "holding qty", "units"]:
        for h in headers:
            if candidate in h:
                mapping["quantity"] = h
                break
        if "quantity" in mapping:
            break

    # Avg price
    for candidate in ["avg", "average", "cost", "buy price"]:
        for h in headers:
            if candidate in h:
                mapping["avg_price"] = h
                break
        if "avg_price" in mapping:
            break

    # Current price
    for candidate in ["ltp", "current", "market price", "close"]:
        for h in headers:
            if candidate in h:
                mapping["current_price"] = h
                break
        if "current_price" in mapping:
            break

    # ISIN
    for h in headers:
        if "isin" in h:
            mapping["isin"] = h
            break

    return ("generic", mapping)


def _detect_mf_broker(headers: list[str]) -> tuple[str, dict]:
    """Detect mutual fund CSV format."""
    normalised = set(headers)
    for broker_id, config in _MF_SIGNATURES.items():
        detect_cols = [c.lower() for c in config["detect"]]
        score = sum(1 for c in detect_cols if c in normalised)
        if score >= 2:
            return (broker_id, config)
    return ("generic_mf", _MF_SIGNATURES.get("generic_mf", {}))


def _find_header_row(rows: list[list]) -> int | None:
    """Find the header row in a list of rows.

    Broker Excel files (e.g. Zerodha) often have metadata/summary rows before
    the actual data table.  The header row is the first row containing
    recognisable column names like 'symbol', 'isin', 'quantity', etc.
    """
    _HEADER_KEYWORDS = {
        "symbol", "isin", "instrument", "scrip", "stock", "company",
        "quantity", "qty", "shares", "units",
        "average price", "avg price", "avg. cost", "avg cost", "avg rate",
        "ltp", "close", "closing price", "current price", "nav",
        "sector", "exchange", "instrument type", "folio", "scheme",
    }

    for i, row in enumerate(rows):
        # Normalise cell values to lowercase strings
        cells = [str(c).strip().lower() for c in row if c is not None]
        if not cells:
            continue
        # Count how many cells match known header keywords
        matches = sum(1 for c in cells if c in _HEADER_KEYWORDS
                      or any(kw in c for kw in _HEADER_KEYWORDS))
        if matches >= 3:
            return i
    return None


def _extract_table_from_sheet(rows: list[list]) -> tuple[list[str], list[dict]]:
    """Given raw rows from an Excel sheet, locate the header row and extract
    the data table as a list of dicts.

    Returns (headers, data_rows).  Headers are lowercase/stripped.
    Empty/None-only rows are skipped.
    """
    header_idx = _find_header_row(rows)
    if header_idx is None:
        return [], []

    raw_headers = rows[header_idx]
    # Build header list, stripping empty columns
    headers: list[str] = []
    col_indices: list[int] = []
    for ci, h in enumerate(raw_headers):
        name = str(h or "").strip().lower()
        if name:
            headers.append(name)
            col_indices.append(ci)

    if not headers:
        return [], []

    data_rows: list[dict] = []
    for row in rows[header_idx + 1:]:
        values = {}
        all_empty = True
        for hi, ci in enumerate(col_indices):
            val = row[ci] if ci < len(row) else None
            s = str(val) if val is not None else ""
            values[headers[hi]] = s
            if s.strip():
                all_empty = False
        if not all_empty:
            data_rows.append(values)

    return headers, data_rows


def _read_all_excel_sheets(wb) -> dict[str, tuple[list[str], list[dict]]]:
    """Read all sheets from an openpyxl Workbook.

    Returns {sheet_name: (headers, data_rows)} for sheets that have data.
    """
    result: dict[str, tuple[list[str], list[dict]]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        # Read all rows including empty cells up to max dimensions
        all_rows: list[list] = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(list(row))
        if not all_rows:
            continue
        headers, data_rows = _extract_table_from_sheet(all_rows)
        if headers and data_rows:
            result[name] = (headers, data_rows)
    return result


def _is_mf_sheet(sheet_name: str, headers: list[str]) -> bool:
    """Heuristic: is this sheet about mutual funds?"""
    name_lower = sheet_name.lower()
    if "mutual" in name_lower or "mf" in name_lower:
        return True
    # Check headers for MF-specific columns
    h_set = set(headers)
    mf_indicators = {"instrument type", "scheme name", "scheme", "folio", "nav", "units"}
    if len(mf_indicators & h_set) >= 2:
        return True
    # Zerodha-style: has 'instrument type' column but no 'sector'
    if "instrument type" in h_set and "sector" not in h_set:
        return True
    return False


def _process_multi_sheet(sheets: dict[str, tuple[list[str], list[dict]]],
                         broker_hint: str = "") -> dict:
    """Process multiple Excel sheets and return combined holdings."""
    all_holdings: list[Holding] = []
    all_mf: list[MFHolding] = []
    errors: list[str] = []
    broker_id = broker_hint or "generic"
    total_rows = 0

    # Prefer 'Combined' sheet if it exists (has both equity + MF)
    # Otherwise process individual sheets
    processed_combined = False
    for sheet_name, (headers, data_rows) in sheets.items():
        if sheet_name.lower() == "combined":
            processed_combined = True

    for sheet_name, (headers, data_rows) in sheets.items():
        # Skip individual sheets if Combined is available
        if processed_combined and sheet_name.lower() not in ("combined",):
            continue

        total_rows += len(data_rows)

        # Detect broker from headers
        if not broker_hint:
            broker_id_detected, mapping = _detect_broker(headers)
            broker_id = broker_id_detected
        else:
            _, mapping = _detect_broker(headers)

        is_mf = _is_mf_sheet(sheet_name, headers)

        for row in data_rows:
            if is_mf or _row_looks_like_mf(row, headers):
                mf = _parse_mf_row(row, headers, mapping, broker_id)
                if mf:
                    all_mf.append(mf)
            else:
                h = _parse_equity_row(row, headers, mapping, broker_id)
                if h:
                    all_holdings.append(h)

    # If Combined sheet mixed equity+MF, separate them (already handled above)
    return {
        "broker": broker_id,
        "holdings": all_holdings,
        "mf_holdings": all_mf,
        "row_count": total_rows,
        "errors": errors,
    }


def _row_looks_like_mf(row: dict, headers: list[str]) -> bool:
    """Check if a row from a combined sheet is a mutual fund row."""
    # Zerodha combined: has 'instrument type' column with values like 'Debt - Liquid', 'Equity - ELSS'
    inst_type = row.get("instrument type", "").strip()
    if inst_type and inst_type != "-":
        return True
    # Check if 'sector' is '-' (MF rows have no sector)
    sector = row.get("sector", "").strip()
    if sector == "-" and inst_type:
        return True
    return False


def _parse_equity_row(row: dict, headers: list[str], mapping: dict,
                      broker_id: str) -> Holding | None:
    """Parse a single row into a Holding using the detected mapping."""
    sym_col = mapping.get("symbol", "")
    isin_col = mapping.get("isin", "")
    name_col = mapping.get("name", "")
    qty_col = mapping.get("quantity", "")
    avg_col = mapping.get("avg_price", "")
    cur_col = mapping.get("current_price", "")
    exch_col = mapping.get("exchange", "")

    symbol = row.get(sym_col, "").strip()
    if not symbol:
        # Fallback: try common column names
        for k in ("symbol", "instrument", "scrip name", "stock code", "script name"):
            if k in row and row[k].strip():
                symbol = row[k].strip()
                break
    if not symbol:
        return None

    qty = _clean_num(row.get(qty_col, ""))
    if qty <= 0:
        # Fallback quantity columns
        for k in headers:
            if "quantity" in k or "qty" in k or "shares" in k:
                qty = _clean_num(row.get(k, ""))
                if qty > 0:
                    break
    if qty <= 0:
        return None

    avg_price = _clean_num(row.get(avg_col, ""))
    if not avg_price:
        for k in headers:
            if "average" in k or "avg" in k or "cost" in k:
                avg_price = _clean_num(row.get(k, ""))
                if avg_price:
                    break

    cur_price = _clean_num(row.get(cur_col, ""))
    if not cur_price:
        for k in headers:
            if "closing" in k or "ltp" in k or "current" in k or "market" in k:
                cur_price = _clean_num(row.get(k, ""))
                if cur_price:
                    break

    isin = row.get(isin_col, "").strip()
    if not isin:
        for k in headers:
            if "isin" in k:
                isin = row.get(k, "").strip()
                break

    return Holding(
        symbol=symbol,
        isin=isin,
        name=row.get(name_col, "").strip() or symbol,
        exchange=row.get(exch_col, "").strip() or "NSE",
        quantity=int(qty) if qty == int(qty) else int(qty),
        avg_price=avg_price,
        current_price=cur_price,
        broker=broker_id,
    )


def _parse_mf_row(row: dict, headers: list[str], mapping: dict,
                   broker_id: str) -> MFHolding | None:
    """Parse a single row into an MFHolding."""
    # Try MF-specific detection
    _, mf_mapping = _detect_mf_broker(headers)

    scheme = row.get(mf_mapping.get("scheme_name", ""), "").strip()
    if not scheme:
        # Fallback: 'symbol' column often has the scheme name
        scheme = row.get("symbol", "").strip()
    if not scheme:
        return None

    units_col = mf_mapping.get("units", "")
    units = _clean_num(row.get(units_col, ""))
    if units <= 0:
        # Fallback: any 'quantity' column
        for k in headers:
            if "quantity" in k or "units" in k:
                units = _clean_num(row.get(k, ""))
                if units > 0:
                    break
    if units <= 0:
        return None

    nav_col = mf_mapping.get("nav", "")
    nav = _clean_num(row.get(nav_col, ""))
    if not nav:
        for k in headers:
            if "closing" in k or "nav" in k or "ltp" in k:
                nav = _clean_num(row.get(k, ""))
                if nav:
                    break

    avg_nav_col = mf_mapping.get("avg_nav", "")
    avg_nav = _clean_num(row.get(avg_nav_col, ""))
    if not avg_nav:
        for k in headers:
            if "average" in k or "avg" in k or "cost" in k:
                avg_nav = _clean_num(row.get(k, ""))
                if avg_nav:
                    break

    invested = avg_nav * units if avg_nav else 0.0
    current = nav * units if nav else 0.0

    isin = ""
    for k in headers:
        if "isin" in k:
            isin = row.get(k, "").strip()
            break

    return MFHolding(
        scheme_code=isin,
        scheme_name=scheme,
        folio=row.get(mf_mapping.get("folio", ""), "").strip(),
        units=units,
        avg_nav=avg_nav,
        current_nav=nav,
        invested_value=round(invested, 2),
        current_value=round(current, 2),
        pnl=round(current - invested, 2),
        pnl_pct=round((current - invested) / invested * 100, 2) if invested else 0.0,
        broker=broker_id,
    )


def import_holdings_csv(
    file_path: str = "",
    content: str = "",
    broker_hint: str = "",
) -> dict:
    """Import holdings from a CSV or Excel file.

    Args:
        file_path: Path to the CSV/XLSX file.
        content: Raw CSV content string (alternative to file_path).
        broker_hint: Optional broker ID hint to skip auto-detection.

    Returns:
        {
            "broker": str,
            "holdings": [Holding, ...],
            "mf_holdings": [MFHolding, ...],
            "row_count": int,
            "errors": [str, ...],
        }
    """
    errors: list[str] = []
    raw_rows: list[dict] = []

    # Read content
    if content:
        reader = csv.DictReader(io.StringIO(content))
        raw_rows = list(reader)
    elif file_path:
        path = Path(file_path)
        if not path.exists():
            return {"broker": "", "holdings": [], "mf_holdings": [],
                    "row_count": 0, "errors": [f"File not found: {file_path}"]}

        suffix = path.suffix.lower()
        if suffix in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, read_only=False, data_only=False)
                all_sheet_rows = _read_all_excel_sheets(wb)
                wb.close()
                if not all_sheet_rows:
                    return {"broker": "", "holdings": [], "mf_holdings": [],
                            "row_count": 0, "errors": ["No data rows found in any sheet"]}
                # Process all sheets and merge results
                return _process_multi_sheet(all_sheet_rows, broker_hint)
            except ImportError:
                return {"broker": "", "holdings": [], "mf_holdings": [],
                        "row_count": 0,
                        "errors": ["openpyxl not installed for Excel support"]}
        else:
            with open(file_path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                raw_rows = list(reader)
    else:
        return {"broker": "", "holdings": [], "mf_holdings": [],
                "row_count": 0, "errors": ["No file or content provided"]}

    if not raw_rows:
        return {"broker": "", "holdings": [], "mf_holdings": [],
                "row_count": 0, "errors": ["No data rows found"]}

    # Normalise headers
    sample = raw_rows[0]
    headers = _normalise_headers(list(sample.keys()))
    normalised_rows = []
    for row in raw_rows:
        normalised_rows.append({k.strip().lower(): v for k, v in row.items()})

    # Detect broker format
    if broker_hint and broker_hint in _BROKER_SIGNATURES:
        broker_id = broker_hint
        mapping = _BROKER_SIGNATURES[broker_hint]
    else:
        broker_id, mapping = _detect_broker(headers)

    # Check if this is a mutual fund CSV
    is_mf = any("scheme" in h or "folio" in h or "nav" in h for h in headers)

    holdings: list[Holding] = []
    mf_holdings: list[MFHolding] = []

    if is_mf:
        _, mf_mapping = _detect_mf_broker(headers)
        for row in normalised_rows:
            scheme = row.get(mf_mapping.get("scheme_name", ""), "").strip()
            if not scheme:
                continue
            units = _clean_num(row.get(mf_mapping.get("units", ""), ""))
            if units <= 0:
                continue
            mf_holdings.append(MFHolding(
                scheme_name=scheme,
                folio=row.get(mf_mapping.get("folio", ""), "").strip(),
                units=units,
                current_nav=_clean_num(row.get(mf_mapping.get("nav", ""), "")),
                current_value=_clean_num(row.get(mf_mapping.get("current", ""), "")),
                invested_value=_clean_num(row.get(mf_mapping.get("invested", ""), "")),
                broker=broker_id,
            ))
    else:
        sym_col = mapping.get("symbol", "")
        isin_col = mapping.get("isin", "")
        name_col = mapping.get("name", "")
        qty_col = mapping.get("quantity", "")
        avg_col = mapping.get("avg_price", "")
        cur_col = mapping.get("current_price", "")
        exch_col = mapping.get("exchange", "")

        for row in normalised_rows:
            symbol = row.get(sym_col, "").strip()
            if not symbol:
                continue
            qty = _clean_num(row.get(qty_col, ""))
            if qty <= 0:
                continue
            holdings.append(Holding(
                symbol=symbol,
                isin=row.get(isin_col, "").strip(),
                name=row.get(name_col, "").strip() or symbol,
                exchange=row.get(exch_col, "").strip() or "NSE",
                quantity=int(qty),
                avg_price=_clean_num(row.get(avg_col, "")),
                current_price=_clean_num(row.get(cur_col, "")),
                broker=broker_id,
            ))

    return {
        "broker": broker_id,
        "holdings": holdings,
        "mf_holdings": mf_holdings,
        "row_count": len(raw_rows),
        "errors": errors,
    }
